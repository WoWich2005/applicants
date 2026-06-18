# -*- coding: utf-8 -*-
"""Строит presentation/Презентация.pptx: берёт копию преза2 (тема, размер 16:9, шрифт Inter),
удаляет все исходные слайды и собирает 11 новых под речь БНТУ в том же визуальном стиле."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy as _copy

FILE = "Презентация.pptx"

# ---------- палитра ----------
BG       = RGBColor(0xF4, 0xF5, 0xF6)
GREEN    = RGBColor(0x00, 0x8A, 0x5E)
GREEN_T  = RGBColor(0xE7, 0xF4, 0xEF)   # светло-зелёная подложка
DARK     = RGBColor(0x20, 0x23, 0x2A)   # заголовки
BODY     = RGBColor(0x3C, 0x3F, 0x45)   # основной текст
MUTED    = RGBColor(0x6B, 0x70, 0x78)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
CARDLINE = RGBColor(0xE6, 0xE8, 0xEB)
FONT     = "Inter"

prs = Presentation(FILE)
SW, SH = prs.slide_width, prs.slide_height            # EMU
IN = 914400
BLANK = prs.slide_layouts[6]                          # «Пустой слайд»

# ---------- удалить все исходные слайды ----------
sldIdLst = prs.slides._sldIdLst
for sld in list(sldIdLst):
    rId = sld.get(qn('r:id'))
    prs.part.drop_rel(rId)
    sldIdLst.remove(sld)

# ---------- низкоуровневые помощники ----------
def soft_shadow(shape, alpha=28000, blur=90000, dist=38000):
    spPr = shape._element.spPr
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)
    eff = spPr.makeelement(qn('a:effectLst'), {})
    sh = eff.makeelement(qn('a:outerShdw'),
                         {'blurRad': str(blur), 'dist': str(dist),
                          'dir': '5400000', 'rotWithShape': '0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val': '8A9099'})
    a = clr.makeelement(qn('a:alpha'), {'val': str(alpha)})
    clr.append(a); sh.append(clr); eff.append(sh); spPr.append(eff)

def no_shadow(shape):
    shape.shadow.inherit = False

def _set_runs(p, text, size, bold, color, align, font=FONT, line=None,
              space_before=None, space_after=None):
    p.alignment = align
    if line is not None:
        p.line_spacing = line
    if space_before is not None:
        p.space_before = Pt(space_before)
    if space_after is not None:
        p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text
    f = r.font
    f.size = Pt(size); f.bold = bold; f.name = font; f.color.rgb = color
    return r

def slide_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG

def new_slide():
    s = prs.slides.add_slide(BLANK)
    slide_bg(s)
    return s

def rrect(slide, x, y, w, h, fill, radius=0.08, line=None, shadow=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    if shadow:
        soft_shadow(sh)
    else:
        no_shadow(sh)
    return sh

def card(slide, x, y, w, h):
    return rrect(slide, x, y, w, h, WHITE, radius=0.06, shadow=True)

def pill(slide, x, y, w, h, text, fill=GREEN, txt=WHITE, size=12.5, bold=True,
         upper=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    sh.adjustments[0] = 0.5
    no_shadow(sh)
    tf = sh.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    tf.margin_left = Pt(8); tf.margin_right = Pt(8)
    _set_runs(tf.paragraphs[0], text.upper() if upper else text, size, bold, txt,
              PP_ALIGN.CENTER)
    return sh

def circle(slide, x, y, d, text, fill=GREEN, txt=WHITE, size=15):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    no_shadow(sh)
    tf = sh.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = 0; tf.margin_bottom = 0
    _set_runs(tf.paragraphs[0], text, size, True, txt, PP_ALIGN.CENTER)
    return sh

def textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, pr in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_runs(p, pr['text'], pr.get('size', 14), pr.get('bold', False),
                  pr.get('color', BODY), pr.get('align', PP_ALIGN.LEFT),
                  line=pr.get('line', 1.08),
                  space_before=pr.get('space_before'),
                  space_after=pr.get('space_after'))
    return tb

def heading(slide, text, size=34):
    textbox(slide, 0.9, 0.55, 14.2, 1.1,
            [{'text': text, 'size': size, 'bold': True, 'color': DARK,
              'align': PP_ALIGN.LEFT, 'line': 1.0}])

def grid(cols, i, left=0.9, total=14.2, gap=0.4):
    w = (total - gap * (cols - 1)) / cols
    return left + i * (w + gap), w

# карточка «заголовок + текст»
def info_card(slide, x, y, w, h, title, body, t_size=17, b_size=13.5,
              t_color=DARK, pad=0.32):
    card(slide, x, y, w, h)
    paras = [{'text': title, 'size': t_size, 'bold': True, 'color': t_color,
              'line': 1.0}]
    if body:
        paras.append({'text': body, 'size': b_size, 'bold': False,
                      'color': BODY, 'line': 1.1, 'space_before': 6})
    textbox(slide, x + pad, y + pad, w - 2 * pad, h - 2 * pad, paras)

# карточка с зелёным бейджем-«таблеткой» сверху
def badge_card(slide, x, y, w, h, badge, body, b_size=13.5):
    card(slide, x, y, w, h)
    pill(slide, x + 0.3, y + 0.28, min(w - 0.6, 0.13 * len(badge) + 0.7), 0.42,
         badge, upper=True, size=11.5)
    textbox(slide, x + 0.32, y + 1.0, w - 0.64, h - 1.2,
            [{'text': body, 'size': b_size, 'color': BODY, 'line': 1.12}])

# ================= СЛАЙД 1 — Титул =================
def s1():
    s = new_slide()
    # правая брендовая панель
    panel = rrect(s, 10.4, 1.5, 4.7, 6.0, GREEN, radius=0.05, shadow=True)
    p_tf = panel.text_frame; p_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_tf.word_wrap = True
    _set_runs(p_tf.paragraphs[0], "BNTU", 54, True, WHITE, PP_ALIGN.CENTER)
    _set_runs(p_tf.add_paragraph(), "Applicants", 40, True, WHITE, PP_ALIGN.CENTER,
              line=1.0)
    _set_runs(p_tf.add_paragraph(), "Система приёмной комиссии", 15, False,
              WHITE, PP_ALIGN.CENTER, line=1.0, space_before=14)
    # бейдж
    pill(s, 0.9, 1.5, 3.1, 0.5, "Дипломный проект", upper=True, size=12.5)
    # тема
    textbox(s, 0.9, 2.25, 9.0, 3.2,
            [{'text': "Система компьютеризации ввода и мониторинга информации "
                      "об абитуриентах при поступлении в БНТУ",
              'size': 31, 'bold': True, 'color': DARK, 'line': 1.1}])
    # исполнитель / руководитель
    textbox(s, 0.9, 6.0, 9.0, 1.6,
            [{'text': "Выполнил: Лемяшевич Владимир Александрович, "
                      "студент группы 10701222", 'size': 15.5, 'color': MUTED,
              'line': 1.25},
             {'text': "Руководитель: Станкевич Сергей Николаевич, "
                      "старший преподаватель кафедры ПОИСиТ", 'size': 15.5,
              'color': MUTED, 'line': 1.25, 'space_before': 8}])

# ================= СЛАЙД 2 — Актуальность =================
def s2():
    s = new_slide()
    heading(s, "Актуальность дипломного проекта")
    data = [
        ("Большие объёмы", "Каждое лето приёмная комиссия обрабатывает тысячи "
         "заявлений по множеству специальностей"),
        ("Несколько специальностей", "Один абитуриент подаёт заявления сразу на "
         "несколько специальностей в порядке предпочтения"),
        ("Постоянные изменения", "Данные всё время меняются, и при каждом "
         "изменении весь конкурс нужно пересчитывать заново"),
        ("Ручная обработка", "Вручную или в обычных таблицах это долго и легко "
         "допустить ошибку"),
    ]
    ys = [2.1, 5.05]; H = 2.7
    for k, (t, b) in enumerate(data):
        x, w = grid(2, k % 2)
        y = ys[k // 2]
        info_card(s, x, y, w, H, t, b)

# ================= СЛАЙД 3 — Существующие решения =================
def s3():
    s = new_slide()
    heading(s, "Существующие решения и их ограничения")
    data = [
        ("Excel-макрос БНТУ", "Написан под один факультет. Чтобы изменить "
         "правила или добавить категорию — нужен программист"),
        ("Зарубежные системы", "Другая модель приёма (решение принимает "
         "сотрудник), высокая цена, чужое законодательство"),
        ("Платформа 1С", "Настроена под российские правила приёма, требует "
         "отдельного специалиста по 1С"),
    ]
    H = 3.3
    for k, (t, b) in enumerate(data):
        x, w = grid(3, k)
        info_card(s, x, 2.05, w, H, t, b, t_size=18)
    # вывод
    bar = rrect(s, 0.9, 5.75, 14.2, 1.55, GREEN_T, radius=0.06, shadow=True)
    pill(s, 1.25, 6.05, 1.7, 0.45, "Вывод", upper=True, size=11.5)
    textbox(s, 3.15, 5.95, 11.7, 1.2,
            [{'text': "Ни одно решение не позволяет настраивать правила отбора "
                      "под белорусскую модель приёма прямо через интерфейс, "
                      "без участия программиста", 'size': 15.5, 'bold': True,
              'color': DARK, 'line': 1.15}],
            anchor=MSO_ANCHOR.MIDDLE)

# ================= СЛАЙД 4 — Цель и задачи =================
def s4():
    s = new_slide()
    heading(s, "Цель и задачи")
    # цель
    rrect(s, 0.9, 1.8, 14.2, 1.5, GREEN_T, radius=0.06, shadow=True)
    pill(s, 1.25, 2.1, 1.6, 0.45, "Цель", upper=True, size=11.5)
    textbox(s, 3.1, 1.98, 11.8, 1.2,
            [{'text': "Разработать веб-приложение, в котором сотрудник приёмной "
                      "комиссии сам задаёт все правила приёма через интерфейс, "
                      "без участия программиста", 'size': 15.5, 'bold': True,
              'color': DARK, 'line': 1.15}], anchor=MSO_ANCHOR.MIDDLE)
    tasks = [
        "Продумать модель данных, в которой правила приёма описываются настройками",
        "Разработать алгоритм автоматического распределения мест по конкурсу",
        "Реализовать разграничение прав и защиту от ошибок ввода данных",
        "Разработать приложение и провести его тестирование",
    ]
    ys = [3.55, 5.5]; H = 1.85
    for k, t in enumerate(tasks):
        x, w = grid(2, k % 2)
        y = ys[k // 2]
        card(s, x, y, w, H)
        circle(s, x + 0.35, y + 0.32, 0.55, str(k + 1))
        textbox(s, x + 1.15, y, w - 1.45, H,
                [{'text': t, 'size': 14.5, 'color': BODY, 'line': 1.12}],
                anchor=MSO_ANCHOR.MIDDLE)

# ================= СЛАЙД 5 — Модель данных =================
def s5():
    s = new_slide()
    heading(s, "Модель данных")
    # цепочка уровней
    chain = ["Факультеты", "Кафедры", "Специальности", "Конкурсные\nсписки",
             "Категории\nприёма"]
    x = 0.9; y = 2.2; bw = 2.55; bh = 1.15; gap = 0.32
    for i, name in enumerate(chain):
        b = rrect(s, x, y, bw, bh, WHITE, radius=0.12, shadow=True)
        tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
        for j, ln in enumerate(name.split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            _set_runs(p, ln, 14.5, True, GREEN, PP_ALIGN.CENTER, line=1.0)
        if i < len(chain) - 1:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                    Inches(x + bw + 0.02), Inches(y + bh/2 - 0.12),
                                    Inches(gap - 0.04), Inches(0.24))
            ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor(0xB9, 0xC0, 0xC7)
            ar.line.fill.background(); no_shadow(ar)
        x += bw + gap
    # карточки-пояснения
    info_card(s, 0.9, 4.05, 6.9, 3.05,
              "Параметры оценки",
              "Перечень показателей, по которым оценивают абитуриентов: баллы "
              "централизованного тестирования, средний балл аттестата, степень "
              "диплома олимпиад и другие. Из них собираются наборы для сравнения.")
    info_card(s, 8.2, 4.05, 6.9, 3.05,
              "Настройка через интерфейс",
              "Структуру, конкурсные списки, категории с квотами и параметры "
              "оценки администратор задаёт сам — без участия программиста. "
              "Абитуриент в системе: его данные, оценки и заявления с приоритетом.")

# ================= СЛАЙД 6 — Алгоритм =================
def s6():
    s = new_slide()
    heading(s, "Автоматическое распределение мест по конкурсу")
    rows = [
        "Абитуриент проходит на самое предпочтительное доступное место из своего списка",
        "Если мест в категории не хватает — остаются сильнейшие, остальные переходят к следующему варианту",
        "Учитывается и квота категории, и общий план приёма на специальность",
        "Пересчёт запускается автоматически при любом изменении данных",
    ]
    y = 2.05; H = 1.15; gap = 0.3
    for k, t in enumerate(rows):
        card(s, 0.9, y, 14.2, H)
        circle(s, 1.25, y + (H - 0.62) / 2, 0.62, str(k + 1))
        textbox(s, 2.25, y, 12.5, H,
                [{'text': t, 'size': 16, 'color': BODY, 'line': 1.1}],
                anchor=MSO_ANCHOR.MIDDLE)
        y += H + gap

# ================= СЛАЙД 7 — Достоверность и роли =================
def s7():
    s = new_slide()
    heading(s, "Достоверность данных и разграничение прав")
    data = [
        ("Журнал действий", "Каждое изменение фиксируется: кто, когда и что "
         "менял. Отдельно ведётся журнал входов в систему"),
        ("Двойной контроль", "Данные, внесённые одним сотрудником, проверяет и "
         "подтверждает другой — это снижает риск ошибки"),
        ("Удаление в два шага", "Один сотрудник запрашивает удаление, другой "
         "подтверждает — случайно потерять данные нельзя"),
        ("Роли пользователей", "У каждой роли — оператор, аудитор, "
         "администратор — свой набор прав на запись"),
    ]
    ys = [2.1, 5.05]; H = 2.7
    for k, (t, b) in enumerate(data):
        x, w = grid(2, k % 2)
        y = ys[k // 2]
        info_card(s, x, y, w, H, t, b)

# ================= СЛАЙД 8 — Демонстрация =================
def s8():
    s = new_slide()
    heading(s, "Демонстрация работы системы")
    # область под видео
    ph = rrect(s, 0.9, 2.05, 9.7, 5.3, RGBColor(0xEA, 0xEC, 0xEF),
               radius=0.04, shadow=True)
    tf = ph.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = True
    _set_runs(tf.paragraphs[0], "▶", 44, True, GREEN, PP_ALIGN.CENTER)
    _set_runs(tf.add_paragraph(), "Видео работы системы", 18, True, MUTED,
              PP_ALIGN.CENTER, line=1.0, space_before=6)
    _set_runs(tf.add_paragraph(), "(вставить .mp4)", 12, False, MUTED,
              PP_ALIGN.CENTER, line=1.0, space_before=2)
    # что на видео
    card(s, 10.9, 2.05, 4.2, 5.3)
    pill(s, 11.2, 2.35, 2.0, 0.42, "На видео", upper=True, size=11.5)
    items = ["Настройка структуры и правил приёма",
             "Создание пользователей и ролей",
             "Ввод абитуриента, оценок и заявлений",
             "Готовый конкурсный список",
             "Авто-пересчёт при изменении данных",
             "Журнал действий и входов",
             "Выгрузка в Excel, смена языка и темы"]
    paras = []
    for i, it in enumerate(items):
        paras.append({'text': "•  " + it, 'size': 12.5, 'color': BODY,
                      'line': 1.1, 'space_before': 0 if i == 0 else 7})
    textbox(s, 11.2, 3.05, 3.6, 4.1, paras)

# ================= СЛАЙД 9 — Технологии =================
def s9():
    s = new_slide()
    heading(s, "Технологии")
    data = [
        ("Клиентская часть", "React", "Одностраничное веб-приложение с "
         "интерфейсом приёмной комиссии"),
        ("Серверная часть", "ASP.NET Core", "REST API: бизнес-логика, "
         "авторизация и расчёт результатов отбора"),
        ("База данных", "PostgreSQL", "Хранение всех данных приёмной кампании "
         "и результатов конкурса"),
    ]
    H = 4.2; y = 2.4
    for k, (badge, tech, body) in enumerate(data):
        x, w = grid(3, k)
        card(s, x, y, w, H)
        pill(s, x + 0.3, y + 0.32, w - 0.6, 0.5, badge, upper=True, size=12)
        textbox(s, x + 0.32, y + 1.1, w - 0.64, H - 1.3,
                [{'text': tech, 'size': 21, 'bold': True, 'color': DARK,
                  'line': 1.0},
                 {'text': body, 'size': 14, 'color': BODY, 'line': 1.15,
                  'space_before': 10}])

# ================= СЛАЙД 10 — Итоги =================
def s10():
    s = new_slide()
    heading(s, "Итоги дипломного проекта")
    data = [
        ("Решение задач", "Создана система ввода и мониторинга данных "
         "абитуриентов с расчётом результатов конкурса"),
        ("Автоматизация", "Конкурс пересчитывается автоматически при любом "
         "изменении исходных данных"),
        ("Гибкость", "Все правила отбора настраиваются через интерфейс, без "
         "участия программиста"),
        ("Надёжность", "Журналирование действий, двойной контроль данных и "
         "разграничение прав пользователей"),
    ]
    ys = [2.1, 5.05]; H = 2.7
    for k, (badge, body) in enumerate(data):
        x, w = grid(2, k % 2)
        y = ys[k // 2]
        badge_card(s, x, y, w, H, badge, body, b_size=14.5)

# ================= СЛАЙД 11 — Спасибо =================
def s11():
    s = new_slide()
    textbox(s, 1.0, 3.6, 14.0, 1.8,
            [{'text': "Спасибо за внимание", 'size': 48, 'bold': True,
              'color': DARK, 'align': PP_ALIGN.CENTER, 'line': 1.0}],
            anchor=MSO_ANCHOR.MIDDLE)

for fn in [s1, s2, s3, s4, s5, s6, s7, s8, s9, s10, s11]:
    fn()

prs.save(FILE)
print("Готово:", FILE, "| слайдов:", len(prs.slides.__iter__.__self__._sldIdLst))
